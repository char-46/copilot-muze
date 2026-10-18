"""Tests for the plugin command registry, skills, docs and scheduler."""

import time

import pytest

from muze.chatlog import ChatLog
from muze.config import MuzeConfig
from muze.core.service import COMMANDS, CommandOutcome, ControlContext, ControlService, Runtime
from muze.docs import html_to_text, read_document
from muze.provider.base import Caps, ProtocolAdapter
from muze.scheduler import Scheduler
from muze.skills import SkillStore
from muze.store.db import open_db
from muze.store.repo import Repository


# ------------------------------------------------------------------ registry
def test_registry_help_scanned_from_declarations():
    text = COMMANDS.help_text()
    assert "/search" in text and "全文搜索" in text
    assert "/schedule" in text and "计划任务" in text
    assert "/reload" in text
    # every declared command appears
    for name in COMMANDS.names():
        assert f"/{name}" in text or name in ("confirm", "cancel")  # merged lines ok


async def test_run_command_guarantees_reply(tmp_path):
    conn = await open_db(str(tmp_path / "r.db"))
    repo = Repository(conn)

    class EchoAdapter(ProtocolAdapter):
        def __init__(self):
            self.caps = Caps(impl_name="t")

        async def start(self):
            raise NotImplementedError

        async def close(self):
            pass

        def events(self):
            raise NotImplementedError

        async def get_login_info(self):
            raise NotImplementedError

        async def get_friend_list(self):
            return []

        async def get_group_list(self):
            return []

        async def get_group_member_name(self, group_id, user_id):
            return ""

        async def send_private_message(self, user_id, segments):
            raise NotImplementedError

        async def send_group_message(self, group_id, segments):
            raise NotImplementedError

        async def get_history_messages(self, scene, peer_id, *, start_message_seq=None, limit=30):
            raise NotImplementedError

        async def get_message(self, scene, peer_id, message_seq):
            return None

        async def get_resource_url(self, resource_id):
            return ""

        async def get_private_file_url(self, user_id, file_id, file_hash):
            return ""

        async def get_group_file_url(self, group_id, file_id):
            return ""

    svc = ControlService(MuzeConfig(), EchoAdapter(), repo, ChatLog(), Runtime())
    ctx = ControlContext(source="cli", chat_key="t", scene="", peer_id=0)

    # unknown -> error reply with help
    out = await svc.run_command("definitely_not_a_cmd")
    assert out.is_error and "可用命令" in out.text

    # local_only enforced for qq ctx
    out = await svc.run_command("restart", ControlContext(source="qq", chat_key="q", scene="friend", peer_id=1))
    assert out.is_error and "本地" in out.text

    # plain string returns are wrapped
    out = await svc.run_command("version")
    assert not out.is_error and out.text.startswith("0.")

    await conn.close()


def test_outcome_chains():
    o = CommandOutcome.chains([[{"type": "text", "data": {"text": "a"}}], [{"type": "text", "data": {"text": "b"}}]])
    assert len(o.extra_chains) == 2


# ------------------------------------------------------------------- skills
def test_skills_parsing_and_override(tmp_path):
    repo_dir = tmp_path / "repo" / "greet"
    user_dir = tmp_path / "user" / "greet"
    repo_dir.mkdir(parents=True)
    user_dir.mkdir(parents=True)
    (repo_dir / "SKILL.md").write_text(
        "---\nname: greet\ndescription: repo version\n---\nREPO BODY", encoding="utf-8"
    )
    (user_dir / "SKILL.md").write_text(
        "---\nname: greet\ndescription: user override\n---\nUSER BODY", encoding="utf-8"
    )
    store = SkillStore([tmp_path / "repo", tmp_path / "user"])
    n = store.reload()
    assert n == 1
    sk = store.get("greet")
    assert sk.body == "USER BODY" and sk.description == "user override"
    assert store.index_lines() == ["- greet: user override"]


def test_skills_no_frontmatter(tmp_path):
    d = tmp_path / "plain"
    d.mkdir()
    (d / "SKILL.md").write_text("just body text", encoding="utf-8")
    store = SkillStore([tmp_path])
    store.reload()
    assert store.get("plain").body == "just body text"


# --------------------------------------------------------------------- docs
def test_html_to_text():
    title, text = html_to_text(
        "<html><head><title>T</title><style>x{}</style></head>"
        "<body><script>bad()</script><p>hello</p><div>world</div></body></html>"
    )
    assert title == "T"
    assert "hello" in text and "world" in text and "bad" not in text


def test_read_document_txt(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("第一行\n第二行", encoding="utf-8")
    assert read_document(p) == "第一行\n第二行"


def test_read_document_docx(tmp_path):
    from docx import Document

    p = tmp_path / "d.docx"
    doc = Document()
    doc.add_paragraph("标题段")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "k1"
    t.cell(0, 1).text = "v1"
    doc.save(str(p))
    text = read_document(p)
    assert "标题段" in text and "k1 | v1" in text


def test_read_document_rejects_doc(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError):
        read_document(p)


# ---------------------------------------------------------------- scheduler
def test_parse_time_spec():
    assert ControlService.parse_time_spec("+3600") == {"delay_sec": 3600}
    assert ControlService.parse_time_spec("3600") == {"delay_sec": 3600}
    assert ControlService.parse_time_spec("every:7200") == {"interval_sec": 7200}
    at = ControlService.parse_time_spec("2030-01-01T09:00:00")
    assert set(at) == {"at_ts"} and at["at_ts"] > int(time.time())
    with pytest.raises(ValueError):
        ControlService.parse_time_spec("明天见")


async def test_scheduler_roundtrip(tmp_path):
    conn = await open_db(str(tmp_path / "s.db"))
    repo = Repository(conn)

    class FakeSvc:
        def __init__(self):
            self.cfg = MuzeConfig()
            self.executed: list[tuple[str, str]] = []
            self.delivered = []

        async def execute(self, prompt, ctx):
            self.executed.append((prompt, ctx.chat_key))
            return CommandOutcome("ok: " + prompt)

        async def deliver(self, ctx, outcome):
            self.delivered.append(outcome.text)

    svc = FakeSvc()
    sch = Scheduler(repo, svc, ChatLog())
    ctx = ControlContext(source="cli", chat_key="t", scene="friend", peer_id=42)

    task = await sch.create(prompt="/status", delay_sec=-0 if False else 0, ctx=ctx)
    assert task["kind"] == "once"
    # due immediately (next_run_at <= now); force tick. tick CLAIMS and
    # spawns concurrent runners — drain() awaits their completion
    n = await sch.tick()
    assert n == 1
    await sch.drain()
    assert svc.executed and svc.executed[0][0] == "/status"
    assert svc.delivered == ["ok: /status"]
    rows = await repo.list_tasks(active_only=False)
    assert rows[0]["status"] == "done" and rows[0]["last_result"].startswith("ok:")
    assert rows[0]["lease_until"] is None

    # interval task reschedules
    t2 = await sch.create(prompt="每日报到", interval_sec=3600, ctx=ctx)
    await sch.tick()  # nothing due yet
    row = await repo.get_task(t2["id"])
    assert row["status"] == "active" and row["next_run_at"] > int(time.time())
    await conn.close()


async def test_scheduler_natural_language_becomes_agent_turn(tmp_path):
    """Regression (live box 2026-10-17 13:11): an agent-made task with a
    natural-language prompt ('重试语音转写：用 transcribe_voice ...') fell
    through execute() as neither-command-nor-:: and the HELP TEXT was
    delivered verbatim into the control chat. Such prompts now get the
    wake prefix prepended; explicit /commands and :: prompts pass through
    untouched."""
    conn = await open_db(str(tmp_path / "s.db"))
    repo = Repository(conn)

    class FakeSvc:
        def __init__(self):
            self.cfg = MuzeConfig()
            self.executed: list[str] = []

        async def execute(self, prompt, ctx):
            self.executed.append(prompt)
            return CommandOutcome("ok")

        async def deliver(self, ctx, outcome):
            pass

    svc = FakeSvc()
    sch = Scheduler(repo, svc, ChatLog())
    ctx = ControlContext(source="cli", chat_key="t", scene="friend", peer_id=42)

    await sch.create(prompt="重试语音转写：用 transcribe_voice 转写 seq 47",
                     delay_sec=0, ctx=ctx)
    await sch.create(prompt="/status", delay_sec=0, ctx=ctx)
    await sch.create(prompt="::直接就是 AI 调用", delay_sec=0, ctx=ctx)
    await sch.tick()
    await sch.drain()  # runners are concurrent; await completion
    assert svc.executed == [
        ":: 重试语音转写：用 transcribe_voice 转写 seq 47",  # prefixed
        "/status",                                            # untouched
        "::直接就是 AI 调用",                                  # untouched
    ]
    await conn.close()


async def test_agent_schedule_task_tool_passes_time_args(tmp_path):
    """Regression (alpha.73): t_schedule_task assembled delay_sec/at_ts/
    interval_sec into kwargs but never passed **kwargs to scheduler.create,
    so every agent-made task died with 'need one of at_ts / delay_sec /
    interval_sec' even with a perfectly valid delay_sec."""
    import json

    from muze.ai.tools import ToolRunner
    from test_permissions import Recorder

    conn = await open_db(str(tmp_path / "st.db"))
    repo = Repository(conn)
    cfg = MuzeConfig()
    cfg.control.controller = [10001]
    service = ControlService(cfg, Recorder(), repo, ChatLog(), Runtime())
    service.runtime.scheduler = Scheduler(repo, service, ChatLog())
    ctx = ControlContext(
        source="qq", chat_key="qq:friend:10001", scene="friend", peer_id=10001
    )
    runner = ToolRunner(service, repo, service.adapter, None, ctx)

    res = json.loads(
        await runner.run(
            "schedule_task",
            {"prompt": "总结 friend/555 最近 50 条", "delay_sec": 3600, "note": "一小时后"},
        )
    )
    assert res.get("created") is True, res
    row = await repo.get_task(res["task_id"])
    assert row["kind"] == "once" and row["status"] == "active"
    assert row["next_run_at"] > int(time.time())

    # at_iso path also reaches the scheduler now
    import datetime as _dt

    iso = (_dt.datetime.now().astimezone() + _dt.timedelta(hours=2)).isoformat()
    res2 = json.loads(
        await runner.run("schedule_task", {"prompt": "到点提醒", "at_iso": iso})
    )
    assert res2.get("created") is True, res2
    await conn.close()
    await conn.close()
